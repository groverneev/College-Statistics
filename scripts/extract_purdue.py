#!/usr/bin/env python3

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any

import openpyxl
import pdfplumber
import xlrd


def normalize_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def parse_number(value: Any) -> float | int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value

    text = str(value).strip()
    if not text:
        return None

    lowered = text.lower()
    if lowered in {"n/a", "na", "not applicable", "varies"}:
        return None

    cleaned = (
        text.replace("$", "")
        .replace(",", "")
        .replace("%", "")
        .replace("\u2212", "-")
        .replace(" ", "")
        .strip()
    )
    if not cleaned:
        return None

    try:
        if re.fullmatch(r"-?\d+", cleaned):
            return int(cleaned)
        return float(cleaned)
    except ValueError:
        return None


def parse_int(value: Any) -> int:
    number = parse_number(value)
    if number is None:
        raise ValueError(f"Unable to parse integer from {value!r}")
    return int(round(float(number)))


def parse_rate(value: Any) -> float:
    number = parse_number(value)
    if number is None:
        return 0.0
    rate = float(number)
    return rate / 100 if rate > 1 else rate


def row_numbers(row: list[Any]) -> list[float | int]:
    numbers: list[float | int] = []
    for cell in row:
        number = parse_number(cell)
        if number is not None:
            numbers.append(number)
    return numbers


def midpoint(p25: int, p75: int) -> int:
    return int(round((p25 + p75) / 2))


def round_half_up(value: float) -> int:
    return int(math.floor(value + 0.5))


def compact_row(values: list[Any]) -> list[Any]:
    compacted: list[Any] = []
    for value in values:
        if value is None:
            continue
        if isinstance(value, str):
            stripped = value.strip()
            if not stripped:
                continue
            compacted.append(stripped)
        else:
            compacted.append(value)
    return compacted


def read_xlsx_rows(path: Path, sheet_name: str) -> list[list[Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        return [compact_row(list(row)) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_xls_rows(path: Path, sheet_name: str) -> list[list[Any]]:
    workbook = xlrd.open_workbook(path)
    sheet_map = {normalize_text(name): name for name in workbook.sheet_names()}
    actual_name = sheet_map.get(normalize_text(sheet_name))
    if actual_name is None:
        raise ValueError(f"Sheet {sheet_name!r} not found in {path.name}")
    sheet = workbook.sheet_by_name(actual_name)

    rows: list[list[Any]] = []
    for row_idx in range(sheet.nrows):
        rows.append(compact_row(sheet.row_values(row_idx)))
    return rows


def read_rows(path: Path, sheet_name: str) -> list[list[Any]]:
    if path.suffix.lower() == ".xlsx":
        return read_xlsx_rows(path, sheet_name)
    if path.suffix.lower() == ".xls":
        return read_xls_rows(path, sheet_name)
    raise ValueError(f"Unsupported workbook type: {path.suffix}")


def find_row(rows: list[list[Any]], *needles: str) -> list[Any]:
    normalized_needles = [normalize_text(needle) for needle in needles]
    for row in rows:
        haystack = normalize_text(" | ".join(str(cell) for cell in row))
        if all(needle in haystack for needle in normalized_needles):
            return row
    raise ValueError(f"Could not find row containing: {needles}")


def find_optional_row(rows: list[list[Any]], *needles: str) -> list[Any] | None:
    try:
        return find_row(rows, *needles)
    except ValueError:
        return None


def find_numeric_row(rows: list[list[Any]], *needles: str) -> list[Any]:
    normalized_needles = [normalize_text(needle) for needle in needles]
    for row in rows:
        haystack = normalize_text(" | ".join(str(cell) for cell in row))
        if all(needle in haystack for needle in normalized_needles) and row_numbers(row):
            return row
    raise ValueError(f"Could not find numeric row containing: {needles}")


def find_optional_numeric_row(rows: list[list[Any]], *needles: str) -> list[Any] | None:
    try:
        return find_numeric_row(rows, *needles)
    except ValueError:
        return None


def extract_last_int(row: list[Any]) -> int:
    numbers = row_numbers(row)
    if not numbers:
        raise ValueError(f"No numeric values found in row: {row}")
    return int(round(float(numbers[-1])))


def extract_excel_admissions(rows: list[list[Any]]) -> dict[str, Any]:
    aggregate_applied = find_optional_numeric_row(rows, "total first-time, first-year students", "applied")
    aggregate_admitted = find_optional_numeric_row(rows, "total first-time, first-year students", "admitted")
    aggregate_enrolled = find_optional_numeric_row(rows, "total first-time, first-year students", "enrolled")

    if aggregate_applied and aggregate_admitted and aggregate_enrolled:
        applied = sum(int(round(float(value))) for value in row_numbers(aggregate_applied))
        admitted = sum(int(round(float(value))) for value in row_numbers(aggregate_admitted))
        enrolled = sum(int(round(float(value))) for value in row_numbers(aggregate_enrolled))
    else:
        applied = sum(
            extract_last_int(find_numeric_row(rows, "total", label, "applied"))
            for label in ("men", "women")
        )
        admitted = sum(
            extract_last_int(find_numeric_row(rows, "total", label, "admitted"))
            for label in ("men", "women")
        )
        enrolled = 0
        for labels in [
            ("total", "enrolled", "full-time", "men"),
            ("total", "enrolled", "part-time", "men"),
            ("total", "enrolled", "full-time", "women"),
            ("total", "enrolled", "part-time", "women"),
            ("total", "enrolled", "full-time", "unknown gender"),
            ("total", "enrolled", "part-time", "unknown gender"),
            ("total", "enrolled", "full-time", "another gender"),
            ("total", "enrolled", "part-time", "another gender"),
        ]:
            row = find_optional_numeric_row(rows, *labels)
            if row is not None:
                enrolled += extract_last_int(row)

    data: dict[str, Any] = {
        "applied": applied,
        "admitted": admitted,
        "enrolled": enrolled,
        "acceptanceRate": round(admitted / applied, 4),
        "yield": round(enrolled / admitted, 4),
    }

    early_applied = find_optional_numeric_row(rows, "number of early decision applications received")
    early_admitted = find_optional_numeric_row(rows, "admitted under early decision")
    if early_applied is not None and early_admitted is not None:
        data["earlyDecision"] = {
            "applied": extract_last_int(early_applied),
            "admitted": extract_last_int(early_admitted),
        }

    return data


def extract_excel_test_scores(rows: list[list[Any]]) -> dict[str, Any]:
    sat_submission_row = find_optional_row(rows, "submitting sat scores") or find_optional_row(rows, "percent submitting sat scores")
    act_submission_row = find_optional_row(rows, "submitting act scores") or find_optional_row(rows, "percent submitting act scores")
    sat_rw_row = find_optional_row(rows, "sat evidence-based reading and writing") or find_optional_row(rows, "sat critical reading")
    sat_math_row = find_row(rows, "sat math")
    sat_composite_row = find_optional_row(rows, "sat composite")
    act_composite_row = find_row(rows, "act composite")

    if sat_rw_row is None:
        raise ValueError("Missing SAT reading/writing row")

    sat_rw_numbers = [int(round(float(value))) for value in row_numbers(sat_rw_row)]
    sat_math_numbers = [int(round(float(value))) for value in row_numbers(sat_math_row)]
    act_numbers = [int(round(float(value))) for value in row_numbers(act_composite_row)]
    sat_composite_numbers = [int(round(float(value))) for value in row_numbers(sat_composite_row or [])]

    if len(sat_composite_numbers) >= 3:
        sat_comp_p25, sat_comp_p50, sat_comp_p75 = sat_composite_numbers[-3:]
    else:
        sat_comp_p25 = sat_rw_numbers[-2] + sat_math_numbers[-2]
        sat_comp_p75 = sat_rw_numbers[-1] + sat_math_numbers[-1]
        sat_comp_p50 = midpoint(sat_comp_p25, sat_comp_p75)

    if len(sat_rw_numbers) >= 3:
        sat_rw_p25, sat_rw_p50, sat_rw_p75 = sat_rw_numbers[-3:]
    else:
        sat_rw_p25, sat_rw_p75 = sat_rw_numbers[-2:]
        sat_rw_p50 = midpoint(sat_rw_p25, sat_rw_p75)

    if len(sat_math_numbers) >= 3:
        sat_math_p25, sat_math_p50, sat_math_p75 = sat_math_numbers[-3:]
    else:
        sat_math_p25, sat_math_p75 = sat_math_numbers[-2:]
        sat_math_p50 = midpoint(sat_math_p25, sat_math_p75)

    if len(act_numbers) >= 3:
        act_p25, act_p50, act_p75 = act_numbers[-3:]
    else:
        act_p25, act_p75 = act_numbers[-2:]
        act_p50 = midpoint(act_p25, act_p75)

    sat_submission_rate = parse_rate(row_numbers(sat_submission_row)[0]) if sat_submission_row else 0.0
    act_submission_rate = parse_rate(row_numbers(act_submission_row)[0]) if act_submission_row else 0.0

    return {
        "sat": {
            "composite": {"p25": sat_comp_p25, "p50": sat_comp_p50, "p75": sat_comp_p75},
            "readingWriting": {"p25": sat_rw_p25, "p50": sat_rw_p50, "p75": sat_rw_p75},
            "math": {"p25": sat_math_p25, "p50": sat_math_p50, "p75": sat_math_p75},
            "submissionRate": round(sat_submission_rate, 4),
        },
        "act": {
            "composite": {"p25": act_p25, "p50": act_p50, "p75": act_p75},
            "submissionRate": round(act_submission_rate, 4),
        },
    }


def extract_excel_demographics(b_rows: list[list[Any]], f_rows: list[list[Any]]) -> dict[str, Any]:
    undergraduate_row = (
        find_optional_numeric_row(b_rows, "total all undergraduates")
        or find_optional_numeric_row(b_rows, "total of all undergraduate students enrolled")
        or find_numeric_row(b_rows, "total undergraduate students")
    )
    graduate_row = (
        find_optional_numeric_row(b_rows, "total all graduate")
        or find_optional_numeric_row(b_rows, "total of all graduate students enrolled")
        or find_numeric_row(b_rows, "total graduate students")
    )
    total_row = find_optional_numeric_row(b_rows, "grand total all students")

    undergraduate = extract_last_int(undergraduate_row)
    graduate = extract_last_int(graduate_row)
    total = extract_last_int(total_row) if total_row is not None else undergraduate + graduate

    by_race = {
        "international": extract_last_int(find_optional_numeric_row(b_rows, "nonresident aliens") or find_numeric_row(b_rows, "nonresidents")),
        "hispanicLatino": extract_last_int(find_numeric_row(b_rows, "hispanic/latino")),
        "blackAfricanAmerican": extract_last_int(find_numeric_row(b_rows, "black or african american, non-hispanic")),
        "white": extract_last_int(find_numeric_row(b_rows, "white, non-hispanic")),
        "asian": extract_last_int(find_numeric_row(b_rows, "asian, non-hispanic")),
        "americanIndianAlaskaNative": extract_last_int(find_numeric_row(b_rows, "american indian or alaska native, non-hispanic")),
        "nativeHawaiianPacificIslander": extract_last_int(find_numeric_row(b_rows, "native hawaiian or other pacific islander, non-hispanic")),
        "twoOrMoreRaces": extract_last_int(find_numeric_row(b_rows, "two or more races, non-hispanic")),
        "unknown": extract_last_int(find_numeric_row(b_rows, "race and/or ethnicity unknown")),
    }

    out_of_state_pct = parse_rate(row_numbers(find_row(f_rows, "percent who are from out of state"))[-1])
    domestic = undergraduate - by_race["international"]
    out_of_state = round_half_up(domestic * out_of_state_pct)
    in_state = domestic - out_of_state

    return {
        "enrollment": {
            "total": total,
            "undergraduate": undergraduate,
            "graduate": graduate,
        },
        "byRace": by_race,
        "byResidency": {
            "inState": in_state,
            "outOfState": out_of_state,
            "international": by_race["international"],
        },
    }


def extract_excel_costs(rows: list[list[Any]]) -> dict[str, int]:
    tuition = extract_last_int(
        find_optional_numeric_row(rows, "in-state (out-of-district)")
        or find_optional_numeric_row(rows, "tuition:", "in-state")
        or find_numeric_row(rows, "in-district")
    )
    fees = extract_last_int(find_optional_numeric_row(rows, "required fees") or find_numeric_row(rows, "fees"))
    room_board_row = find_optional_numeric_row(rows, "food and housing") or find_optional_numeric_row(rows, "room and board")
    if room_board_row is None:
        raise ValueError("Missing room and board row")
    room_and_board = extract_last_int(room_board_row)

    return {
        "tuition": tuition,
        "fees": fees,
        "roomAndBoard": room_and_board,
        "totalCOA": tuition + fees + room_and_board,
    }


def extract_excel_financial_aid(rows: list[list[Any]]) -> dict[str, Any]:
    total_students = extract_last_int(find_numeric_row(rows, "number of degree-seeking undergraduate students"))
    awarded_aid = extract_last_int(find_numeric_row(rows, "awarded any financial aid"))
    fully_met_rate = parse_rate(find_numeric_row(rows, "percentage of need that was met")[-1])
    avg_package = extract_last_int(find_numeric_row(rows, "average financial aid package"))
    avg_grant = extract_last_int(
        find_optional_numeric_row(rows, "average need-based scholarship and grant award")
        or find_optional_numeric_row(rows, "average need-based scholarship or grant award")
        or find_numeric_row(rows, "average need-based grant award")
    )

    return {
        "percentReceivingAid": round(awarded_aid / total_students, 4),
        "averageAidPackage": avg_package,
        "averageNeedBasedGrant": avg_grant,
        "percentNeedFullyMet": round(fully_met_rate, 4),
    }


def extract_excel_year_data(workbook_path: Path) -> dict[str, Any]:
    admissions_rows = read_rows(workbook_path, "CDS-C")
    demographics_rows = read_rows(workbook_path, "CDS-B")
    student_life_rows = read_rows(workbook_path, "CDS-F")
    costs_rows = read_rows(workbook_path, "CDS-G")
    aid_rows = read_rows(workbook_path, "CDS-H")

    return {
        "admissions": extract_excel_admissions(admissions_rows),
        "testScores": extract_excel_test_scores(admissions_rows),
        "demographics": extract_excel_demographics(demographics_rows, student_life_rows),
        "costs": extract_excel_costs(costs_rows),
        "financialAid": extract_excel_financial_aid(aid_rows),
    }


class PurduePdfExtractor:
    def __init__(self, pdf_path: Path):
        self.pdf_path = pdf_path
        self.pages: list[str] = []
        self.tables: list[list[list[str | None]]] = []
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                self.pages.append(page.extract_text() or "")
                for table in page.extract_tables():
                    self.tables.append(table)
        self.text = "\n".join(self.pages)

    def search(self, patterns: list[str], flags: int = re.IGNORECASE | re.MULTILINE | re.DOTALL) -> re.Match[str]:
        for pattern in patterns:
            match = re.search(pattern, self.text, flags)
            if match:
                return match
        raise ValueError(f"No pattern matched for {self.pdf_path.name}: {patterns[0]}")

    def search_optional(self, patterns: list[str], flags: int = re.IGNORECASE | re.MULTILINE | re.DOTALL) -> re.Match[str] | None:
        for pattern in patterns:
            match = re.search(pattern, self.text, flags)
            if match:
                return match
        return None

    def _match_int(self, patterns: list[str]) -> int:
        return parse_int(self.search(patterns).group(1))

    def _find_table_row(self, label: str) -> list[str | None]:
        normalized = normalize_text(label)
        for table in self.tables:
            for row in table:
                if not row:
                    continue
                first_cell = normalize_text(row[0] if row[0] is not None else "")
                row_text = normalize_text(" | ".join("" if cell is None else str(cell) for cell in row))
                if first_cell == normalized or row_text.startswith(normalized):
                    return row
        raise ValueError(f"Missing table row {label!r} in {self.pdf_path.name}")

    def _find_table(self, header: str) -> list[list[str | None]]:
        normalized = normalize_text(header)
        for table in self.tables:
            if not table or not table[0]:
                continue
            first_row = normalize_text(" | ".join("" if cell is None else str(cell) for cell in table[0]))
            if normalized in first_row:
                return table
        raise ValueError(f"Missing table with header {header!r} in {self.pdf_path.name}")

    def _scores_from_pattern(self, patterns: list[str]) -> tuple[int, int, int]:
        match = self.search(patterns)
        numbers = [parse_int(group) for group in match.groups() if group is not None]
        if len(numbers) == 3:
            return numbers[0], numbers[1], numbers[2]
        if len(numbers) == 2:
            return numbers[0], midpoint(numbers[0], numbers[1]), numbers[1]
        raise ValueError(f"Unexpected score match in {self.pdf_path.name}: {match.groups()}")

    def extract_admissions(self) -> dict[str, Any]:
        total_applied_match = self.search_optional([
            r"Total first-time, first-year \(degree-seeking\) who applied\s+([\d,\s]+)",
        ])
        total_admitted_match = self.search_optional([
            r"Total first-time, first-year \(degree-seeking\) who were admitted\s+([\d,\s]+)",
        ])
        total_enrolled_match = self.search_optional([
            r"Total first-time, first-year \(degree-seeking\) who enrolled\s+([\d,\s]+)",
        ])
        male_applied_match = self.search_optional([
            r"Total first-time, first-year males who applied\s+([\d,\s]+)",
            r"Total first-time, first-year men who applied\s+([\d,\s]+)",
            r"Total first-time, first-year \(freshman\) men who applied\s+([\d,\s]+)",
        ])
        female_applied_match = self.search_optional([
            r"Total first-time, first-year females who applied\s+([\d,\s]+)",
            r"Total first-time, first-year women who applied\s+([\d,\s]+)",
            r"Total first-time, first-year \(freshman\) women who applied\s+([\d,\s]+)",
        ])
        male_admitted_match = self.search_optional([
            r"Total first-time, first-year males who were admitted\s+([\d,\s]+)",
            r"Total first-time, first-year men who were admitted\s+([\d,\s]+)",
            r"Total first-time, first-year \(freshman\) men who were admitted\s+([\d,\s]+)",
        ])
        female_admitted_match = self.search_optional([
            r"Total first-time, first-year females who were admitted\s+([\d,\s]+)",
            r"Total first-time, first-year women who were admitted\s+([\d,\s]+)",
            r"Total first-time, first-year \(freshman\) women who were admitted\s+([\d,\s]+)",
        ])
        male_enrolled_match = self.search_optional([
            r"Total full-time, first-time, first-year males who enrolled\s+([\d,\s]+)",
            r"Total full-time, first-time, first-year men who enrolled\s+([\d,\s]+)",
            r"Total full-time, first-time, first-year \(freshman\) men who enrolled\s+([\d,\s]+)",
        ])
        male_part_time_match = self.search_optional([
            r"Total part-time, first-time, first-year males who enrolled\s+([\d,\s]+)",
            r"Total part-time, first-time, first-year men who enrolled\s+([\d,\s]+)",
            r"Total part-time, first-time, first-year \(freshman\) men who enrolled\s+([\d,\s]+)",
        ])
        female_enrolled_match = self.search_optional([
            r"Total full-time, first-time, first-year females who enrolled\s+([\d,\s]+)",
            r"Total full-time, first-time, first-year women who enrolled\s+([\d,\s]+)",
            r"Total full-time, first-time, first-year \(freshman\) women who enrolled\s+([\d,\s]+)",
        ])
        female_part_time_match = self.search_optional([
            r"Total part-time, first-time, first-year females who enrolled\s+([\d,\s]+)",
            r"Total part-time, first-time, first-year women who enrolled\s+([\d,\s]+)",
            r"Total part-time, first-time, first-year \(freshman\) women who enrolled\s+([\d,\s]+)",
        ])

        if total_applied_match and total_admitted_match and total_enrolled_match:
            applied = parse_int(total_applied_match.group(1))
            admitted = parse_int(total_admitted_match.group(1))
            enrolled = parse_int(total_enrolled_match.group(1))
        elif male_applied_match and female_applied_match:
            applied = parse_int(male_applied_match.group(1)) + parse_int(female_applied_match.group(1))
            admitted = parse_int(male_admitted_match.group(1)) + parse_int(female_admitted_match.group(1))
            enrolled = (
                parse_int(male_enrolled_match.group(1))
                + parse_int(male_part_time_match.group(1))
                + parse_int(female_enrolled_match.group(1))
                + parse_int(female_part_time_match.group(1))
            )
        else:
            applied_match = self.search([r"Total first-time, first-year students who applied in Fall \d{4}\s+([\d,\.]+)\s+([\d,\.]+)"])
            admitted_match = self.search([r"Total first-time, first-year students admitted in Fall \d{4}\s+([\d,\.]+)\s+([\d,\.]+)"])
            full_time_match = self.search([r"Full-time, first-time, first-year students enrolled in Fall \d{4}\s+([\d,\.]+)\s+([\d,\.]+)"])
            part_time_match = self.search([r"Part-time, first-time, first-year students enrolled in Fall \d{4}\s+([\d,\.]+)\s+([\d,\.]+)"])
            applied = parse_int(applied_match.group(1)) + parse_int(applied_match.group(2))
            admitted = parse_int(admitted_match.group(1)) + parse_int(admitted_match.group(2))
            enrolled = (
                parse_int(full_time_match.group(1))
                + parse_int(full_time_match.group(2))
                + parse_int(part_time_match.group(1))
                + parse_int(part_time_match.group(2))
            )

        return {
            "applied": applied,
            "admitted": admitted,
            "enrolled": enrolled,
            "acceptanceRate": round(admitted / applied, 4),
            "yield": round(enrolled / admitted, 4),
        }

    def extract_test_scores(self) -> dict[str, Any]:
        sat_submission = parse_rate(
            self.search([
                r"Submitting SAT Scores\s+(\d+(?:\.\d+)?)%",
                r"Percent Number\s+Submitting SAT Scores\s+(\d+(?:\.\d+)?)%",
            ]).group(1)
        )
        act_submission = parse_rate(
            self.search([
                r"Submitting ACT Scores\s+(\d+(?:\.\d+)?)%",
                r"Percent Number\s+Submitting ACT Scores\s+(\d+(?:\.\d+)?)%",
            ]).group(1)
        )

        parsed_table: list[list[str | None]] | None = None
        try:
            parsed_table = self._find_table("Assessment")
        except ValueError:
            parsed_table = None

        if parsed_table is not None:
            score_rows: dict[str, list[str | None]] = {}
            for row in parsed_table[1:]:
                if not row or not row[0]:
                    continue
                score_rows[normalize_text(row[0])] = row

            def row_scores(label_options: list[str]) -> tuple[int, int, int]:
                for label in label_options:
                    row = score_rows.get(normalize_text(label))
                    if row:
                        values = [parse_int(cell) for cell in row[1:] if parse_number(cell) is not None]
                        if len(values) == 3:
                            return values[0], values[1], values[2]
                        if len(values) == 2:
                            return values[0], midpoint(values[0], values[1]), values[1]
                raise ValueError(f"Missing score row {label_options[0]!r} in {self.pdf_path.name}")

            sat_rw = row_scores(["SAT Evidence-Based Reading and Writing", "SAT Evidence-Based Reading and", "SAT Critical Reading"])
            sat_math = row_scores(["SAT Math"])
            act = row_scores(["ACT Composite"])

            sat_composite_row = score_rows.get(normalize_text("SAT Composite"))
            if sat_composite_row:
                composite_numbers = [parse_int(cell) for cell in sat_composite_row[1:] if parse_number(cell) is not None]
                if len(composite_numbers) == 3:
                    sat_composite = (composite_numbers[0], composite_numbers[1], composite_numbers[2])
                elif len(composite_numbers) == 2:
                    sat_composite = (composite_numbers[0], midpoint(composite_numbers[0], composite_numbers[1]), composite_numbers[1])
                else:
                    sat_composite = (
                        sat_rw[0] + sat_math[0],
                        sat_rw[1] + sat_math[1],
                        sat_rw[2] + sat_math[2],
                    )
            else:
                sat_composite = (
                    sat_rw[0] + sat_math[0],
                    sat_rw[1] + sat_math[1],
                    sat_rw[2] + sat_math[2],
                )
        else:
            sat_composite = self._scores_from_pattern([r"SAT Composite .*?(\d{3,4})\s+(\d{3,4})\s+(\d{3,4})"])
            sat_rw = self._scores_from_pattern([
                r"SAT Evidence-Based Reading and Writing .*?(\d{3})\s+(\d{3})\s+(\d{3})",
                r"SAT Evidence-Based Reading and\s+Writing .*?(\d{3})\s+(\d{3})\s+(\d{3})",
            ])
            sat_math = self._scores_from_pattern([r"SAT Math .*?(\d{3})\s+(\d{3})\s+(\d{3})"])
            act = self._scores_from_pattern([r"ACT Composite .*?(\d{2})\s+(\d{2})\s+(\d{2})"])

        return {
            "sat": {
                "composite": {"p25": sat_composite[0], "p50": sat_composite[1], "p75": sat_composite[2]},
                "readingWriting": {"p25": sat_rw[0], "p50": sat_rw[1], "p75": sat_rw[2]},
                "math": {"p25": sat_math[0], "p50": sat_math[1], "p75": sat_math[2]},
                "submissionRate": round(sat_submission, 4),
            },
            "act": {
                "composite": {"p25": act[0], "p50": act[1], "p75": act[2]},
                "submissionRate": round(act_submission, 4),
            },
        }

    def extract_demographics(self) -> dict[str, Any]:
        undergraduate = self._match_int([
            r"Total all undergraduates\s+([\d,\.]+)",
            r"Total of all undergraduate students enrolled\s+([\d,\.]+)",
        ])
        graduate = self._match_int([
            r"Total all graduate\s+([\d,\.]+)",
            r"Total of all graduate students enrolled\s+([\d,\.]+)",
        ])
        total_match = self.search_optional([r"GRAND TOTAL ALL STUDENTS\s+([\d,\.]+)"])
        total = parse_int(total_match.group(1)) if total_match else undergraduate + graduate

        by_race = {
            "international": self._match_int([
                r"International \(nonresidents\)\s+[\d,\.]+\s+[\d,\.]+\s+([\d,\.]+)",
                r"Nonresidents?\s+[\d,\.]+\s+[\d,\.]+\s+([\d,\.]+)",
                r"Nonresident aliens\s+[\d,\.]+\s+[\d,\.]+\s+([\d,\.]+)",
            ]),
            "hispanicLatino": self._match_int([r"Hispanic/Latino\s+[\d,\.]+\s+[\d,\.]+\s+([\d,\.]+)"]),
            "blackAfricanAmerican": self._match_int([r"Black or African American,\s*non-?\s*Hispanic.*?(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)"]),
            "white": self._match_int([r"White,\s*non-?\s*Hispanic.*?(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)"]),
            "asian": self._match_int([r"Asian,\s*non-?\s*Hispanic.*?(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)"]),
            "americanIndianAlaskaNative": self._match_int([
                r"American Indian or Alaska Native,\s*non-?\s*Hispanic.*?(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)",
                r"American Indian or Alaska Native,\s*non-\s*(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)\s*Hispanic",
            ]),
            "nativeHawaiianPacificIslander": self._match_int([
                r"Native Hawaiian or other Pacific Islander,\s*non-?\s*Hispanic.*?(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)",
                r"Native Hawaiian or other Pacific Islander,\s*non-?Hispa.*?(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)",
                r"Native Hawaiian or other Pacific Islander,\s*(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)\s*non-?\s*Hispanic",
            ]),
            "twoOrMoreRaces": self._match_int([r"Two or more races,\s*non-?\s*Hispanic.*?(?:[\d,\.]+)\s+(?:[\d,\.]+)\s+([\d,\.]+)"]),
            "unknown": self._match_int([r"Race and/or ethnicity unknown\s+[\d,\.]+\s+[\d,\.]+\s+([\d,\.]+)"]),
        }

        residency_match = self.search([
            r"Percent who are from out of state \(exclude.*?(\d+(?:\.\d+)?)%\s+(\d+(?:\.\d+)?)%",
            r"Percent who are from out of state \(exclude international/\s*non-?residents from the numerator and denominator\)\s+(\d+(?:\.\d+)?)%\s+(\d+(?:\.\d+)?)%",
        ])
        out_of_state_pct = parse_rate(residency_match.group(2))
        domestic = undergraduate - by_race["international"]
        out_of_state = round_half_up(domestic * out_of_state_pct)
        in_state = domestic - out_of_state

        return {
            "enrollment": {
                "total": total,
                "undergraduate": undergraduate,
                "graduate": graduate,
            },
            "byRace": by_race,
            "byResidency": {
                "inState": in_state,
                "outOfState": out_of_state,
                "international": by_race["international"],
            },
        }

    def extract_costs(self) -> dict[str, int]:
        # Starting in 2024-2025 Purdue splits section G by campus. The UI only supports
        # one COA series, so use the flagship West Lafayette section for continuity.
        cost_text = self.text
        west_lafayette_match = self.search_optional([
            r"G\. ANNUAL EXPENSES - Purdue West Lafayette(?P<section>.*?)(?:G\. ANNUAL EXPENSES - Purdue in Indianapolis|H\.)",
        ])
        if west_lafayette_match:
            cost_text = west_lafayette_match.group("section")

        tuition = parse_int(
            self.search_optional(
                [r"Tuition: In-state \(out-of-district\): \$?([\d,]+(?:\.\d+)?)"],
                flags=re.IGNORECASE | re.MULTILINE | re.DOTALL,
            ).group(1)
            if self.search_optional([r"Tuition: In-state \(out-of-district\): \$?([\d,]+(?:\.\d+)?)"], flags=re.IGNORECASE | re.MULTILINE | re.DOTALL)
            else self.search([r"Tuition:\s+\$?([\d,]+(?:\.\d+)?)"], flags=re.IGNORECASE | re.MULTILINE | re.DOTALL).group(1)
        )
        fees_match = re.search(r"Required Fees:\s*\$?([\d,]+(?:\.\d+)?)", cost_text, re.IGNORECASE | re.MULTILINE | re.DOTALL)
        room_match = re.search(
            r"(?:Food and housing \(on-campus\)|Room and Board(?:\s*\(on-campus\))?)\s*:\s*\$?([\d,]+(?:\.\d+)?)",
            cost_text,
            re.IGNORECASE | re.MULTILINE | re.DOTALL,
        )
        if room_match is None:
            raise ValueError(f"Missing room and board value in {self.pdf_path.name}")

        fees = parse_int(fees_match.group(1)) if fees_match else 0
        room_and_board = parse_int(room_match.group(1))

        return {
            "tuition": tuition,
            "fees": fees,
            "roomAndBoard": room_and_board,
            "totalCOA": tuition + fees + room_and_board,
        }

    def extract_financial_aid(self) -> dict[str, Any]:
        h2_table: list[list[str | None]] | None = None
        hk_table: list[list[str | None]] | None = None
        for table in self.tables:
            if len(table) > 1 and len(table[1]) > 1:
                row_text = normalize_text(" | ".join("" if cell is None else str(cell) for cell in table[1]))
                if "number of degree-seeking undergraduate students" in row_text:
                    h2_table = table
            if table and table[0] and table[0][0] in {"H", "I", "J", "K"}:
                hk_table = table

        if h2_table is not None:
            h2_rows = {normalize_text(row[0]): row for row in h2_table[1:] if row and row[0]}
            total_students = parse_int(h2_rows["a"][3])
            awarded_aid = parse_int(h2_rows["d"][3])
        else:
            total_students = parse_int(
                self.search([
                    r"Number of degree-seeking undergraduate\s+students \(CDS Item B1 if reporting on Fall \d{4}\s+([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)\s+cohort\)",
                    r"Number of degree-seeking undergraduate students.*?([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)\s+cohort",
                    r"Number of degree-seeking undergraduate students.*?([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)",
                ]).group(2)
            )
            awarded_aid = parse_int(
                self.search([
                    r"Number of students in line (?:c|\(C\)) who were awarded any financial\s*([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)\s*aid",
                    r"Number of students in line (?:c|\(C\)) who were awarded\s*([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)\s*any financial aid",
                    r"Number of students in line (?:c|\(C\)) who were awarded any financial aid\s*([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)",
                ]).group(2)
            )

        if hk_table is not None:
            hk_rows = {normalize_text(row[0]): row for row in hk_table if row and row[0]}
            percent_need_fully_met = parse_rate(hk_rows["i"][3]) if "i" in hk_rows else parse_rate(
                self.search([
                    r"On average, the percentage of need that was met.*?([\d\.]+%)\s+([\d\.]+%)\s+([\d\.]+%)",
                ]).group(2)
            )
            average_aid_package = parse_int(hk_rows["j"][3]) if "j" in hk_rows else parse_int(
                self.search([
                    r"The average financial aid package of those in line.*?\$ ?([\d,\.]+)\s+\$ ?([\d,\.]+)\s+\$ ?([\d,\.]+)",
                ]).group(2)
            )
            average_need_based_grant = parse_int(hk_rows["k"][3]) if "k" in hk_rows else parse_int(
                self.search([
                    r"Average need-based scholarship(?: and grant| or grant)? award.*?\$ ?([\d,\.]+)\s+\$ ?([\d,\.]+)\s+\$ ?([\d,\.]+)",
                ]).group(2)
            )
        else:
            percent_need_fully_met = parse_rate(
                self.search([
                    r"On average, the percentage of need that was met.*?([\d\.]+%)\s+([\d\.]+%)\s+([\d\.]+%)",
                ]).group(2)
            )
            average_aid_package = parse_int(
                self.search([
                    r"The average financial aid package of those in line.*?\$ ?([\d,\.]+)\s+\$ ?([\d,\.]+)\s+\$ ?([\d,\.]+)",
                ]).group(2)
            )
            average_need_based_grant = parse_int(
                self.search([
                    r"Average need-based scholarship(?: and grant| or grant)? award.*?\$ ?([\d,\.]+)\s+\$ ?([\d,\.]+)\s+\$ ?([\d,\.]+)",
                ]).group(2)
            )

        return {
            "percentReceivingAid": round(awarded_aid / total_students, 4),
            "averageAidPackage": average_aid_package,
            "averageNeedBasedGrant": average_need_based_grant,
            "percentNeedFullyMet": round(percent_need_fully_met, 4),
        }

    def extract_all(self) -> dict[str, Any]:
        return {
            "admissions": self.extract_admissions(),
            "testScores": self.extract_test_scores(),
            "demographics": self.extract_demographics(),
            "costs": self.extract_costs(),
            "financialAid": self.extract_financial_aid(),
        }


def year_label(path: Path) -> str:
    match = re.search(r"(20\d{2})[-_](20\d{2})", path.stem)
    if not match:
        raise ValueError(f"Could not determine year from filename: {path.name}")
    return f"{match.group(1)}-{match.group(2)}"


def validate_year(year: str, year_data: dict[str, Any]) -> None:
    undergrad = year_data["demographics"]["enrollment"]["undergraduate"]
    race_total = sum(year_data["demographics"]["byRace"].values())
    residency_total = sum(year_data["demographics"]["byResidency"].values())
    if race_total != undergrad:
        raise ValueError(f"{year}: race total {race_total} does not match undergrad {undergrad}")
    if residency_total != undergrad:
        raise ValueError(f"{year}: residency total {residency_total} does not match undergrad {undergrad}")


def build_dataset(input_dir: Path) -> dict[str, Any]:
    years: dict[str, Any] = {}
    for path in sorted(input_dir.iterdir()):
        if path.suffix.lower() not in {".pdf", ".xls", ".xlsx"}:
            continue
        year = year_label(path)
        if path.suffix.lower() in {".xls", ".xlsx"}:
            years[year] = extract_excel_year_data(path)
        else:
            years[year] = PurduePdfExtractor(path).extract_all()
        validate_year(year, years[year])

    return {
        "name": "Purdue University",
        "slug": "purdue",
        "years": years,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract Purdue CDS data from mixed Excel and PDF sources.")
    parser.add_argument("--input-dir", type=Path, default=Path("College-Data/Purdue"))
    parser.add_argument("--output", type=Path, default=Path("src/data/schools/purdue.json"))
    args = parser.parse_args()

    dataset = build_dataset(args.input_dir)
    args.output.write_text(json.dumps(dataset, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
